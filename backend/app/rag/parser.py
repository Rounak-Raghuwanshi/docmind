"""Turn uploaded bytes into pages of text.

PDF pages map 1:1 to pages. DOCX / TXT / Markdown have no pages, so they are split into
sections (by heading, capped at ~SECTION_WORDS words) that play the role of pages for
citations.
"""

import io
import re
import zipfile
from collections.abc import Callable, Iterator

from app.rag.ocr import ocr_png, tesseract_available
from app.rag.types import Page

SECTION_WORDS = 500

PDF = "application/pdf"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TXT = "text/plain"
MD = "text/markdown"

EXTENSIONS = {".pdf": PDF, ".docx": DOCX, ".txt": TXT, ".md": MD, ".markdown": MD}


class IngestionError(Exception):
    """A permanent, user-readable failure (retrying won't help)."""


def detect_content_type(filename: str, head: bytes, data: bytes | None = None) -> str:
    """Trust magic bytes, not the extension or the client's Content-Type header."""
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    expected = EXTENSIONS.get(ext)
    if expected is None:
        raise IngestionError("Unsupported file type. Upload PDF, DOCX, TXT or Markdown.")

    if expected == PDF:
        if not head.startswith(b"%PDF-"):
            raise IngestionError("File has a .pdf extension but is not a PDF")
        return PDF
    if expected == DOCX:
        if not head.startswith(b"PK\x03\x04"):
            raise IngestionError("File has a .docx extension but is not a Word document")
        if data is not None:
            try:
                with zipfile.ZipFile(io.BytesIO(data)) as zf:
                    if "word/document.xml" not in zf.namelist():
                        raise IngestionError("File is a zip archive, not a Word document")
            except zipfile.BadZipFile as e:
                raise IngestionError("Word document is corrupted") from e
        return DOCX
    # Text formats: must be valid UTF-8 without NUL bytes (binary files fail here).
    sample = data if data is not None else head
    if b"\x00" in sample:
        raise IngestionError("File is binary, not text")
    try:
        sample.decode("utf-8")
    except UnicodeDecodeError as e:
        if data is not None or len(sample) < 4:  # a truncated head may split a code point
            raise IngestionError("Text files must be UTF-8 encoded") from e
    return expected


def count_pages(content_type: str, data: bytes) -> int:
    if content_type == PDF:
        import pymupdf

        with pymupdf.open(stream=data, filetype="pdf") as doc:
            return int(doc.page_count)
    return len(_sections_for(content_type, data))


def iter_pages(
    content_type: str,
    data: bytes,
    *,
    ocr_enabled: bool = True,
    ocr_min_chars: int = 20,
    ocr_dpi: int = 200,
) -> Iterator[Page]:
    if content_type == PDF:
        yield from _iter_pdf(data, ocr_enabled, ocr_min_chars, ocr_dpi)
    else:
        for i, text in enumerate(_sections_for(content_type, data), start=1):
            yield Page(number=i, text=text)


def _iter_pdf(data: bytes, ocr_enabled: bool, min_chars: int, dpi: int) -> Iterator[Page]:
    import pymupdf

    try:
        doc = pymupdf.open(stream=data, filetype="pdf")
    except Exception as e:  # PyMuPDF raises several types for broken files
        raise IngestionError("PDF is corrupted or unreadable") from e
    with doc:
        if doc.needs_pass:
            raise IngestionError("PDF is password-protected")
        can_ocr = ocr_enabled and tesseract_available()
        for page in doc:  # type: ignore[attr-defined]
            # "blocks" keeps paragraph structure; sort=True gives natural reading order.
            blocks = page.get_text("blocks", sort=True)
            text = "\n\n".join(b[4].strip() for b in blocks if b[6] == 0 and b[4].strip())
            if len(text.strip()) < min_chars and can_ocr:
                png = page.get_pixmap(dpi=dpi).tobytes("png")
                yield Page(number=page.number + 1, text=ocr_png(png), ocr=True)
            else:
                yield Page(number=page.number + 1, text=text)


def _sections_for(content_type: str, data: bytes) -> list[str]:
    if content_type == DOCX:
        return _docx_sections(data)
    text = data.decode("utf-8", errors="replace")
    is_heading: Callable[[str], bool]
    if content_type == MD:
        is_heading = lambda line: bool(re.match(r"^#{1,6}\s", line))  # noqa: E731
    else:
        is_heading = lambda line: False  # noqa: E731
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    return _group_sections(paragraphs, is_heading)


def _docx_sections(data: bytes) -> list[str]:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as e:
        raise IngestionError("Word document is corrupted or unreadable") from e

    paragraphs: list[str] = []
    headings: set[int] = set()
    for p in document.paragraphs:
        text = p.text.strip()
        if not text:
            continue
        style = (p.style.name if p.style is not None else "") or ""
        if style.lower().startswith(("heading", "title")):
            headings.add(len(paragraphs))
            text = "# " + text  # normalise to Markdown so the chunker recognises headings
        paragraphs.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                paragraphs.append(" | ".join(cells))
    return _group_sections(paragraphs, lambda line: line.startswith("# "))


def _group_sections(paragraphs: list[str], is_heading: Callable[[str], bool]) -> list[str]:
    """Start a new section at each heading or when the current one passes SECTION_WORDS."""
    sections: list[list[str]] = []
    current: list[str] = []
    words = 0
    for para in paragraphs:
        n = len(para.split())
        if current and (is_heading(para) or words + n > SECTION_WORDS):
            sections.append(current)
            current, words = [], 0
        current.append(para)
        words += n
    if current:
        sections.append(current)
    return ["\n\n".join(s) for s in sections]
