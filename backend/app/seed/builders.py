"""Render demo documents as real files, so seeding goes through the normal ingestion pipeline."""

import html
import io

from app.seed.documents import DemoDoc

CSS = """
body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #1e293b; }
h1 { font-size: 20pt; color: #312e81; margin-bottom: 6pt; }
h2 { font-size: 14pt; color: #3730a3; margin-top: 14pt; margin-bottom: 4pt; }
p { margin-bottom: 6pt; }
li { margin-bottom: 3pt; }
"""


def _html(doc: DemoDoc) -> str:
    parts = [f"<h1>{html.escape(doc.title)}</h1>"]
    for s in doc.sections:
        parts.append(f"<h2>{html.escape(s.heading)}</h2>")
        parts += [f"<p>{html.escape(p)}</p>" for p in s.paragraphs]
        if s.bullets:
            parts.append(
                "<ul>" + "".join(f"<li>{html.escape(b)}</li>" for b in s.bullets) + "</ul>"
            )
    return "".join(parts)


def build_pdf(doc: DemoDoc) -> bytes:
    import pymupdf

    story = pymupdf.Story(html=_html(doc), user_css=CSS)
    page_rect = pymupdf.paper_rect("a4")
    content = page_rect + (60, 70, -60, -70)  # noqa: RUF005 - Rect inset, not a tuple
    buf = io.BytesIO()
    writer = pymupdf.DocumentWriter(buf)
    more = True
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content)
        story.draw(device)
        writer.end_page()
    writer.close()

    # Running header and footer on every page: the cleaner should strip these.
    pdf = pymupdf.open(stream=buf.getvalue(), filetype="pdf")
    for page in pdf:  # type: ignore[attr-defined]
        page.insert_text(
            (60, 40), f"DocMind demo · {doc.title}", fontsize=8, color=(0.45, 0.45, 0.5)
        )
        page.insert_text(
            (60, page_rect.height - 30),
            f"Page {page.number + 1} of {pdf.page_count}",
            fontsize=8,
            color=(0.45, 0.45, 0.5),
        )
    data = pdf.tobytes(garbage=3, deflate=True)
    pdf.close()
    if doc.scanned:
        return _as_scan(data)
    return bytes(data)


def _as_scan(pdf_bytes: bytes) -> bytes:
    """Rasterise every page so the PDF has no text layer, like a scanned paper document."""
    import pymupdf

    src = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    out = pymupdf.open()
    for page in src:  # type: ignore[attr-defined]
        pix = page.get_pixmap(dpi=150, colorspace=pymupdf.csGRAY)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, pixmap=pix)
    data = out.tobytes(deflate=True)
    src.close()
    out.close()
    return bytes(data)


def build_docx(doc: DemoDoc) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading(doc.title, level=0)
    for s in doc.sections:
        d.add_heading(s.heading, level=1)
        for p in s.paragraphs:
            d.add_paragraph(p)
        for b in s.bullets:
            d.add_paragraph(b, style="List Bullet")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def build_markdown(doc: DemoDoc) -> bytes:
    lines = [f"# {doc.title}", ""]
    for s in doc.sections:
        lines += [f"## {s.heading}", ""]
        for p in s.paragraphs:
            lines += [p, ""]
        lines += [f"- {b}" for b in s.bullets]
        if s.bullets:
            lines.append("")
    return "\n".join(lines).encode()


CONTENT_TYPES = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "md": "text/markdown",
}


def build(doc: DemoDoc) -> tuple[bytes, str]:
    data = {"pdf": build_pdf, "docx": build_docx, "md": build_markdown}[doc.kind](doc)
    return data, CONTENT_TYPES[doc.kind]
